
"""Turn FMS CSV (often a nameless GUID file) into a real .xlsx WPS can open."""

from __future__ import annotations

import csv
import io
import os
from datetime import datetime
from pathlib import Path

from lib.wait_file import classify_bytes


from lib import cli
info = cli.info
ok = cli.ok
warn = cli.warn
error = cli.error


def decode_text(data: bytes) -> str:
    if data.startswith(b"\xff\xfe") or data.startswith(b"\xfe\xff"):
        return data.decode("utf-16")
    if data.startswith(b"\xef\xbb\xbf"):
        return data.decode("utf-8-sig")
    for enc in ("utf-8-sig", "utf-8", "gb18030", "cp1252"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def detect_delim(first_line: str) -> str:
    tabs = first_line.count("\t")
    commas = first_line.count(",")
    semis = first_line.count(";")
    if tabs >= 2 and tabs >= commas:
        return "\t"
    if commas >= 2 and commas >= semis:
        return ","
    if semis >= 2:
        return ";"
    return "\t"


def parse_fms_text(text: str) -> list[list[str]]:
    lines = [ln for ln in text.splitlines() if ln.strip()]
    if not lines:
        return []
    delim = detect_delim(lines[0])
    info("FMS delimiter:", "TAB" if delim == "\t" else delim)
    if delim == "\t":
        return [ln.split("\t") for ln in lines]
    return [list(r) for r in csv.reader(io.StringIO(text), delimiter=delim)]


def csv_bytes_to_xlsx(data: bytes, dest: Path) -> Path:
    from openpyxl import Workbook

    rows = parse_fms_text(decode_text(data))
    if not rows:
        raise ValueError("CSV has no rows")
    ncols = max(len(r) for r in rows)
    info("Header:", " | ".join(rows[0][:6]), ("…" if len(rows[0]) > 6 else ""))
    if len(rows) >= 2 and rows[1]:
        info("First Order ID:", rows[1][0])
        headers = [h.strip().lower() for h in rows[0]]
        try:
            di = headers.index("destination station")
            info("First Destination Station:", rows[1][di] if di < len(rows[1]) else "")
        except ValueError:
            pass
    if ncols < 3:
        raise ValueError(
            f"Only {ncols} column(s) — file was not split. Do not open this in WPS/MC."
        )

    wb = Workbook()
    ws = wb.active
    ws.title = "Reception"
    for r_i, row in enumerate(rows, 1):
        for c_i, val in enumerate(row, 1):
            cell = ws.cell(r_i, c_i, "" if val is None else str(val))
            cell.number_format = "@"
    dest.parent.mkdir(exist_ok=True)
    wb.save(dest)
    ok(f"Wrote {len(rows)} rows × {ncols} cols → {dest}")
    return dest


def xlsx_to_tsv(path: Path) -> str:
    import shutil
    import tempfile

    from openpyxl import load_workbook

    src = path
    tmp = Path(tempfile.gettempdir()) / f"mc_{path.name}"
    try:
        shutil.copy2(path, tmp)
        src = tmp
    except OSError:
        tsv = path.with_suffix(".tsv")
        if tsv.is_file():
            info("xlsx locked by WPS — using", tsv)
            return tsv.read_text(encoding="utf-8-sig", errors="replace")

    wb = load_workbook(src, read_only=True, data_only=True)
    ws = wb.active
    lines: list[str] = []
    for row in ws.iter_rows(values_only=True):
        if all(v is None or str(v).strip() == "" for v in row):
            continue
        lines.append("\t".join("" if v is None else str(v) for v in row))
    wb.close()
    if not lines:
        raise ValueError(f"No rows in {path}")
    info(f"TSV for MC: {len(lines)} rows from {path.name}")
    return "\n".join(lines)


def to_real_xlsx(raw: Path, dest_dir: Path, stamp: str | None = None) -> Path:
    dest_dir.mkdir(exist_ok=True)
    stamp = stamp or datetime.now().strftime("%Y%m%d_%H%M%S")
    if "playwright-artifacts" in str(raw).lower() or "e5f9e491" in raw.name.lower():
        raise ValueError(
            f"Refusing stale download {raw}. "
            "Wait for export_reception_order_*.csv in Downloads."
        )
    data = raw.read_bytes()
    ext, kind = classify_bytes(data)
    info(f"Source: {raw} ({len(data)} bytes)  {kind}")

    tsv_copy = dest_dir / f"FMS_Reception_{stamp}.tsv"
    tsv_copy.write_bytes(data)

    out = dest_dir / f"FMS_Reception_{stamp}.xlsx"
    if kind.startswith("Excel (.xlsx)"):
        out.write_bytes(data)
    elif kind.startswith("CSV") or ext in {".csv", ".tsv", ".bin"} or kind == "unknown":
        csv_bytes_to_xlsx(data, out)
    elif kind.startswith("Excel 97"):
        raise ValueError(
            "This is a fake/old .xls (CSV saved with the wrong name). "
            "Use the nameless ~320 KB file or the .csv, not the .xls."
        )
    else:
        raise ValueError(f"Cannot turn {kind} into xlsx")

    user = Path.home() / "Downloads" / out.name
    if user.resolve() != out.resolve():
        user.write_bytes(out.read_bytes())
    info("WPS file:", user)
    return user


def save_original_csv(raw: Path, dest_dir: Path, stamp: str | None = None) -> Path:
    """Keep FMS bytes as-is (tab CSV). Do not convert dates or columns."""
    dest_dir.mkdir(exist_ok=True)
    stamp = stamp or datetime.now().strftime("%Y%m%d_%H%M%S")
    data = raw.read_bytes()
    text = decode_text(data)
    lines = [ln for ln in text.splitlines() if ln.strip()]
    info(f"Source: {raw} ({len(data)} bytes)")
    info(f"Rows including header: {len(lines)}")
    if lines:
        info("Header starts:", lines[0][:80])
    if len(lines) >= 2:
        first = lines[1].split("\t")[0] if "\t" in lines[1] else lines[1][:40]
        info("First Order ID:", first)
    sample = text[:12000]
    if "Calamba City" in sample and "Davao" not in sample:
        info("WARNING: this looks like an OLD reverse-logistics file, not the Davao batch.")
        info("Do not use it for WPS / the next website.")
    out = dest_dir / f"FMS_Reception_{stamp}.csv"
    out.write_bytes(data)
    user = Path.home() / "Downloads" / out.name
    if user.resolve() != out.resolve():
        user.write_bytes(data)
    info("CSV for WPS / next website:", user)
    return user


def open_in_wps(path: Path) -> bool:
    """Launch the export in WPS and make sure the window actually comes up.

    Returns True when the window is up and in front, False otherwise — so
    the caller can stop claiming "Opened in WPS" when it was not.

    os.startfile() alone was the whole of this function, and that is why the
    export only appeared once you clicked the terminal. The launch works;
    the window activation does not. Windows only lets a window take the
    foreground if the calling process owns the input queue, and a console
    app running the pipeline does not — WPS came up behind the terminal and
    stayed there.

    So: grant foreground rights first (ASFW_ANY, which survives WPS's own
    launcher), launch, then wait for the window and force it in front with
    AttachThreadInput + the ALT unlock.

    The wait still only matches VISIBLE windows, deliberately. WPS leaves
    a hidden window behind for every document it closes (it holds the file
    lock and keeps the title), so matching hidden windows too would find
    last run's export and bring the wrong file forward. Measured: after
    closing a WPS document, its titled window stays alive with
    IsWindowVisible false while WPS shows a blank "Workbook1".
    """
    info("Opening in WPS / default app:", path)
    try:
        from lib import wps_copy
    except ImportError as e:
        warn("WPS helpers unavailable, opening blind:", e)
        os.startfile(path)  # Windows
        return False

    try:
        wps_copy.allow_launch_to_take_focus()
    except Exception as e:
        warn("Could not grant foreground rights to WPS:", e)

    os.startfile(path)  # Windows

    try:
        hwnd, title = wps_copy.wait_for_export_window(path, timeout=60)
    except TimeoutError as e:
        # WPS queues a new open behind a "Save the Document?" modal it keeps
        # from a window closed without answering (manual close, or an earlier
        # session). Esc dismisses it and the queued launch then shows its
        # window — try once before giving up.
        try:
            front = wps_copy.front_window_title()
            if "Save the Document" in front or "保存" in front:
                wps_copy.dismiss_dialog()
                hwnd, title = wps_copy.wait_for_export_window(path, timeout=15)
                if hwnd and wps_copy.bring_to_front(hwnd):
                    ok("WPS window is in front (after dismissing a "
                       "modal):", title)
                    return True
        except Exception:
            pass
        warn("Could not auto-open WPS:", e)
        info("Open this .xlsx yourself:", path)
        return False

    if wps_copy.bring_to_front(hwnd):
        ok("WPS window is in front:", title)
        return True
    warn("WPS opened but is NOT in front — click its taskbar icon.")
    warn("Foreground is:", wps_copy.front_window_title() or "(none)")
    return False


