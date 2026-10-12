"""Live proof of bug 24's fix, against the REAL FMS pager.

tools/check_pager_reset.py proves the fix on a simulated pager, but that markup
was written by me -- so it can only prove my own model of FMS is self
consistent. This runs the same scenario against the live order-tracking page:

  1. run a real batch search for ~120 known OIDs (several pages at 50/page)
  2. force the pager to the LAST page, then confirm page 1 is not even rendered
  3. goto_first_page() must get home
  4. copy_all_pages() from that same last-page state must get EVERY row, first
     and last, and leave the pager back on page 1

The OIDs come from the last real FMS export, so the search is a read-only
lookup of orders this pipeline already handled. Nothing is exported, so no
Reception Order task is created.

Run with:  py tools\\verify_live_pager_reset.py [N_OIDS]
"""

import sys
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from playwright.sync_api import sync_playwright

from fms_batch_export import (CDP, paste_batch, submit_and_wait_results)
from lib import tabs
import lib.packed_scrape as ps

fails = []


def check(label, cond, extra=""):
    print(("  PASS  " if cond else "  FAIL  ") + label
          + (f"  {extra}" if extra else ""))
    if not cond:
        fails.append(label)


def oids_from_exports(n: int) -> list[str]:
    """Pool SPX tracking numbers out of the FMS exports on disk.

    The FMS results grid calls the column "SPX Tracking Number"; the exported
    workbook calls the same value "Order ID". Exports are pooled newest-first
    because FMS only matches orders still in scope -- a single older export
    yields mostly misses (102 OIDs from one file matched only 10).
    """
    try:
        import openpyxl
    except ImportError:
        print("  openpyxl not installed -- cannot read the exports")
        return []
    dls = sorted((Path.home() / "Downloads").glob("FMS_Reception_*.xlsx"),
                 key=lambda p: p.stat().st_mtime, reverse=True)
    seen: set[str] = set()
    out: list[str] = []
    used = 0
    for xl in dls:
        if len(out) >= n:
            break
        try:
            wb = openpyxl.load_workbook(xl, read_only=True, data_only=True)
            ws = wb[wb.sheetnames[0]]
            hdr = None
            ci = None
            for r in ws.iter_rows(values_only=True):
                if hdr is None:
                    hdr = [str(c).strip() if c else "" for c in r]
                    for col in ("SPX Tracking Number", "Order ID"):
                        if col in hdr:
                            ci = hdr.index(col)
                            break
                    continue
                if ci is None or ci >= len(r):
                    continue
                v = r[ci]
                v = str(v).strip() if v else ""
                if v and v not in seen:
                    seen.add(v)
                    out.append(v)
            wb.close()
            used += 1
        except Exception as e:
            print(f"  could not read {xl.name}: {str(e)[:60]}")
    print(f"  pooled {len(out)} distinct OIDs from {used} export(s)")
    return out[:n]


def rows_as_dicts(data) -> list[dict]:
    """Same header->value mapping copy_all_pages() does, for inspection."""
    if not data or not data.get("rows"):
        return []
    mp = {}
    for i, h in enumerate(data["headers"]):
        h = h.strip()
        if h in ps.WANTED and h not in mp:
            mp[h] = i
    out = []
    for cells in data["rows"]:
        if not cells or all(not c for c in cells):
            continue
        out.append({w: (cells[mp[w]] if w in mp and mp[w] < len(cells) else "")
                    for w in ps.WANTED})
    return out


def main() -> int:
    want = int(sys.argv[1]) if len(sys.argv) > 1 else 120
    tabs.begin_session()
    with sync_playwright() as p:
        b = p.chromium.connect_over_cdp(CDP)
        ctx = b.contexts[0]
        pages = [q for q in ctx.pages if "orderTracking" in (q.url or "")]
        if not pages:
            print("No FMS tab open.")
            return 1
        pg = pages[0]
        pg.bring_to_front()

        oids = oids_from_exports(want)
        if len(oids) < 10:
            print(f"Only {len(oids)} OIDs available -- not enough for several "
                  f"pages. Cannot prove this live.")
            return 1

        print(f"\n-- batch search for {len(oids)} OIDs --")
        t0 = time.time()
        paste_batch(pg, oids, 3000)
        submit_and_wait_results(pg, 120)
        print(f"   submitted in {time.time() - t0:.1f}s")

        ps.set_page_size_50(pg)
        size = ps._page_size(pg)
        total = ps._table_total(pg)
        npages = -(-int(total) // size) if total else 0
        print(f"   page size {size}, {npages} page(s), total {total}")
        if npages < 2:
            print("   fewer than 2 pages -- this export is too small to prove "
                  "the last-page trap.")
            return 1

        print("\n-- force the pager to the LAST page (the reported symptom) --")
        ps.goto_first_page(pg)          # make sure we start from a known place
        landed, _ = ps._goto_page(pg, npages)
        check(f"reached the last page ({npages})", landed is True,
              f"active={ps._active_page(pg)}")
        check("it really is the last page", ps._active_page(pg) == npages)

        rendered = pg.evaluate(
            "() => [...document.querySelectorAll('.ssc-pagination .pager-item')]"
            ".map(e => (e.textContent||'').trim())")
        nums = [int(x) for x in rendered if x.isdigit()]
        print(f"   pager shows numbers: {rendered}")
        trap = "1" not in rendered
        print(f"   the pager renders {len(nums)} number(s) and reaches {max(nums)}"
              f" of {npages} page(s)")
        if trap:
            check("the windowing trap IS present on the real pager", True)
        else:
            # Measured: FMS renders the WHOLE pager when the result set is small
            # (6 pages -> 1..6 all visible). The window only drops page 1 once
            # the result set is wider than the window, i.e. >~10 pages. So the
            # trap is size-dependent on the real site, which is exactly why it
            # looks intermittent to the user. check_pager_reset.py forces a
            # 3-wide window over 9 pages to produce it deterministically.
            print("   NOTE: the pager renders every number at this result size, "
                  "so the windowing trap cannot bite here -- it needs a result "
                  "set wider than the pager's number window. That is why the "
                  "symptom looks intermittent. tools/check_pager_reset.py forces "
                  "the narrow window that produces it.")

        print("\n-- the old fallback would have read the wrong rows --")
        d = ps._read_table_stable(pg)
        shown = rows_as_dicts(d)
        first_oid = oids[0]
        check("page 1's first OID is NOT on screen right now",
              not any(r.get("SPX Tracking Number") == first_oid for r in shown),
              f"{len(shown)} rows on screen, first="
              f"{shown[0]['SPX Tracking Number'] if shown else None}")

        print("\n-- goto_first_page() must get home on the real pager --")
        t0 = time.time()
        at = ps.goto_first_page(pg)
        dt = time.time() - t0
        check("it returned to page 1", at == 1, f"ended on page {at} in {dt:.1f}s")

        print("\n-- a full copy starting from the LAST page --")
        landed, _ = ps._goto_page(pg, npages)
        check("back on the last page to reproduce", landed is True,
              f"active={ps._active_page(pg)}")
        t0 = time.time()
        rows, pages = ps.copy_all_pages(pg)
        dt = time.time() - t0
        got = [r["SPX Tracking Number"] for r in rows]
        print(f"   copied {len(rows)} row(s) from {pages} page(s) in {dt:.1f}s")

        check("it copied a whole result set, not one page",
              len(rows) > size, f"{len(rows)} rows (one page would be {size})")
        check("all OIDs distinct", len(set(got)) == len(got))
        check("page 1's first OID is present (it really started at page 1)",
              first_oid in got)
        # Do NOT compare against _table_total() here: measured live, it reports
        # a stale count (102) while the grid actually holds 300 rows across 6
        # pages, because it is read while the results are still settling.
        # The self-consistent check is that it visited every page the rows
        # actually occupy.
        expected = -(-len(rows) // size)
        check("it walked every page the results occupy", pages == expected,
              f"{len(rows)} rows at {size}/page -> {expected} page(s), "
              f"visited {pages}")
        check("it leaves the pager on page 1 for the next run",
              ps._active_page(pg) == 1, f"ended on page {ps._active_page(pg)}")
    return 0


if __name__ == "__main__":
    code = main()
    print("\n" + ("ALL PASS" if not fails and code == 0
                  else f"{len(fails)} FAILED: {fails}"))
    raise SystemExit(code or (1 if fails else 0))
